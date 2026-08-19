from io import BytesIO
from pathlib import Path
import re
from typing import Any

from docx import Document as DocxDocument
from pypdf import PdfReader


class UnsupportedDocument(ValueError):
    pass


def _paragraph_spans(text: str, page: int | None = None) -> list[dict[str, Any]]:
    spans = []
    offset = 0
    section = None
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            offset += len(raw) + 1
            continue
        if line.startswith("#"):
            section = line.lstrip("# ")
        spans.append({"index": len(spans), "page": page, "section": section, "text": line, "start": offset, "end": offset + len(raw)})
        offset += len(raw) + 1
    return spans


def parse_document(filename: str, content: bytes) -> list[dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        return _paragraph_spans(content.decode("utf-8"))
    if suffix == ".pdf":
        spans: list[dict[str, Any]] = []
        for page_number, page in enumerate(PdfReader(BytesIO(content)).pages, start=1):
            # Layout extraction preserves the vertical gaps between authored
            # paragraphs. Rejoin visual line wraps so provenance is a meaningful
            # paragraph/row rather than an arbitrary PDF glyph line.
            layout = page.extract_text(extraction_mode="layout") or ""
            section = None
            for block in re.split(r"\n\s*\n", layout):
                text = " ".join(block.split())
                if not text:
                    continue
                if len(text.split()) <= 8 and not text.endswith((".", ":", ";")):
                    section = text
                spans.append({"index": len(spans), "page": page_number, "section": section, "text": text, "start": None, "end": None})
        return spans
    if suffix == ".docx":
        spans = []
        section = None
        for paragraph in DocxDocument(BytesIO(content)).paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                section = text
            spans.append({"index": len(spans), "page": None, "section": section, "text": text, "start": None, "end": None})
        return spans
    raise UnsupportedDocument("Only PDF, DOCX, and TXT documents are supported")
