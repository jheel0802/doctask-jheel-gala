from io import BytesIO

from docx import Document
from reportlab.pdfgen.canvas import Canvas

from project_evidence.parsing import parse_document


def test_txt_parser_has_stable_evidence_span():
    spans = parse_document("notes.txt", b"# Planning\nLaunch remains delayed.\n")
    assert spans[1] == {"index": 1, "page": None, "section": "Planning", "text": "Launch remains delayed.", "start": 11, "end": 34}


def test_docx_parser_preserves_section():
    document = Document()
    document.add_heading("Risks", level=1)
    document.add_paragraph("Vendor capacity is constrained.")
    buffer = BytesIO()
    document.save(buffer)
    spans = parse_document("risk.docx", buffer.getvalue())
    assert spans[-1]["section"] == "Risks"
    assert spans[-1]["text"] == "Vendor capacity is constrained."


def test_pdf_parser_preserves_page():
    buffer = BytesIO()
    canvas = Canvas(buffer)
    canvas.drawString(72, 720, "Launch remains delayed.")
    canvas.save()
    spans = parse_document("status.pdf", buffer.getvalue())
    assert spans[0]["page"] == 1
    assert "Launch remains delayed" in spans[0]["text"]
